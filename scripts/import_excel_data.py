"""Import supplied PSSPL workbooks into local Tally snapshot tables."""
import io, re, sqlite3
from datetime import date, datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "crm.sqlite"
MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
def workbook(name): return openpyxl.load_workbook(io.BytesIO((ROOT / name).read_bytes()), data_only=True, read_only=True)
def number(value):
    try: return float(value or 0)
    except (TypeError, ValueError): return 0.0
def text(value): return str(value).strip() if value is not None else ""
def iso(value): return value.strftime("%Y-%m-%d") if isinstance(value, (date, datetime)) else (text(value) or None)
def slug(value): return re.sub(r"[^A-Z0-9]+", "-", text(value).upper()).strip("-")

DB.parent.mkdir(exist_ok=True)
con = sqlite3.connect(DB); con.execute("PRAGMA foreign_keys=ON")
tables = ["tally_customer_monthly_sales", "tally_product_monthly_sales", "tally_rm_product_sales", "tally_rm_targets", "tally_party_payments", "tally_salespeople", "tally_customers", "tally_products"]
for table in tables: con.execute(f"DELETE FROM {table}")

source = "6) Dealer Summary Apr, 2023 to May, 2024.xlsx"
periods = ["2023-04", "2023-05", "2023-06", "2023-07", "2023-08", "2023-09", "2023-10", "2023-11", "2023-12", "2024-01", "2024-02", "2024-03", "2024-04", "2024-05"]
for row in workbook(source)["Master Data"].iter_rows(min_row=5, values_only=True):
    if not text(row[1]): continue
    customer_id = text(row[2]).split("/")[0].strip() or f"PARTY-{slug(row[1])}"
    con.execute("INSERT OR REPLACE INTO tally_customers VALUES (?,?,?,?,?,?,?,?,?,?,?)", (customer_id, text(row[1]), text(row[3]), text(row[4]), text(row[5]), text(row[6]), text(row[7]), number(row[8]), number(row[23]), iso(row[24]), source))
    for period, amount in zip(periods, row[9:23]): con.execute("INSERT OR REPLACE INTO tally_customer_monthly_sales VALUES (?,?,?)", (customer_id, period, number(amount)))

source = "8) PSS Product Analysis I 2022-23 (005).xlsx"
for row in workbook(source)["Original Summary "].iter_rows(min_row=5, values_only=True):
    if not text(row[2]): continue
    product_id = text(row[1]) or f"PRODUCT-{slug(row[2])}"; monthly = row[7:19]
    total_qty = number(row[19]) or sum(number(v) for v in monthly); rate = number(row[6])
    con.execute("INSERT OR REPLACE INTO tally_products VALUES (?,?,?,?,?,?,?,?,?)", (product_id, text(row[2]), text(row[3]), text(row[4]), rate, total_qty, rate * total_qty, sum(number(v) != 0 for v in monthly), source))
    for idx, qty in enumerate(monthly):
        year, month = (2022, idx + 4) if idx < 9 else (2023, idx - 8)
        con.execute("INSERT OR REPLACE INTO tally_product_monthly_sales VALUES (?,?,?)", (product_id, f"{year}-{month:02d}", number(qty)))

source = "10) COMPARISON OF  PRODUCT WITH QTY.xlsx"
for ws in workbook(source).worksheets:
    start = 4 if ws.title == "Anil Jena" else 3; aggregates = {}
    for row in ws.iter_rows(min_row=start, values_only=True):
        product = text(row[1])
        if not product or product.lower().startswith("total"): continue
        for fy, qty, sales in (("2024-25", row[2], row[4]), ("2025-26", row[3], row[5])):
            old = aggregates.get((product, fy), (0, 0)); aggregates[(product, fy)] = (old[0] + number(qty), old[1] + number(sales))
    for (product, fy), values in aggregates.items(): con.execute("INSERT OR REPLACE INTO tally_rm_product_sales VALUES (?,?,?,?,?)", (ws.title.strip(), product, fy, values[0], values[1]))

wb = workbook("9) PSS Sales Incentive Calculations-24-25.xlsx")
for row in wb["Targets"].iter_rows(min_row=4, values_only=True):
    rm = text(row[3])
    if not rm: continue
    for idx, target in enumerate(row[5:17]):
        year, month = (2024, idx + 4) if idx < 9 else (2025, idx - 8)
        con.execute("INSERT OR REPLACE INTO tally_rm_targets VALUES (?,?,?)", (rm, f"{year}-{month:02d}", number(target)))

current_rm = None
for row in wb["Sales Exec Sales"].iter_rows(min_row=5, values_only=True):
    code, name = text(row[2]), " ".join(text(row[3]).split())
    if not name: continue
    if number(row[1]) == 0: current_rm = name; continue
    if code: con.execute("INSERT OR REPLACE INTO tally_salespeople VALUES (?,?,?,?,?)", (code, name, current_rm, "", iso(row[5])))

current_rm = None
for row in workbook("RM name with sales person..xlsx")["Sheet1"].iter_rows(min_row=2, values_only=True):
    if text(row[1]): current_rm = text(row[1])
    person, city = text(row[3]), text(row[4])
    if not person: continue
    match = con.execute("SELECT code FROM tally_salespeople WHERE lower(replace(name,' ',''))=lower(replace(?,' ',''))", (person,)).fetchone()
    if match: con.execute("UPDATE tally_salespeople SET rm_name=?,city=? WHERE code=?", (current_rm, city, match[0]))
    else: con.execute("INSERT OR REPLACE INTO tally_salespeople VALUES (?,?,?,?,?)", (f"PERSON-{slug(person)}", person, current_rm, city, None))

for row in workbook("2) All Party wise Payments.xlsx")["Party wise Payments"].iter_rows(min_row=15, max_row=15, values_only=True):
    for idx, amount in enumerate(row[2:14]):
        year, month = (2024, idx + 4) if idx < 9 else (2025, idx - 8)
        con.execute("INSERT OR REPLACE INTO tally_party_payments VALUES (?,?,?)", ("2 Daughters", f"{year}-{month:02d}", number(amount)))

stamp = datetime.now().isoformat(timespec="seconds")
con.execute("INSERT OR REPLACE INTO tally_import_metadata VALUES ('last_import_at',?)", (stamp,)); con.execute("INSERT OR REPLACE INTO tally_import_metadata VALUES ('source','PSSPL supplied Excel workbooks')")
con.commit(); print({"imported_at": stamp, **{t: con.execute(f'SELECT COUNT(*) FROM {t}').fetchone()[0] for t in tables}}); con.close()
